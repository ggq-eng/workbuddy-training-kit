# -*- coding: utf-8 -*-
"""中文公文式 Word 排版工具库（封面 / 目录域 / 编号章节 / 带边框表格 / 页眉页脚）"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------- 配色 ----------
NAVY = RGBColor(0x1F, 0x3B, 0x63)      # 主深蓝
NAVY_HEX = "1F3B63"
BLUE = RGBColor(0x1F, 0x4E, 0x79)      # 次级蓝
BLUE_HEX = "1F4E79"
HEAD_FILL = "DCE6F1"                    # 表头浅蓝底纹
ALT_FILL = "F5F7FA"                     # 隔行浅底
GREY = RGBColor(0x59, 0x59, 0x59)

CN_BODY = "宋体"
CN_HEAD = "黑体"
EN_BODY = "Times New Roman"


def set_run(run, cn=CN_BODY, size=12, bold=False, color=None, italic=False):
    run.font.name = EN_BODY
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), cn)
    rFonts.set(qn('w:ascii'), EN_BODY)
    rFonts.set(qn('w:hAnsi'), EN_BODY)
    return run


def new_document(margin_cm=(2.6, 2.6, 2.5, 2.5)):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = EN_BODY
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), CN_BODY)
    sec = doc.sections[0]
    sec.left_margin = Cm(margin_cm[0])
    sec.right_margin = Cm(margin_cm[1])
    sec.top_margin = Cm(margin_cm[2])
    sec.bottom_margin = Cm(margin_cm[3])
    pf = st.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.4
    pf.space_after = Pt(4)
    return doc


def set_para(p, align=None, first_indent=0, before=0, after=4, spacing=1.4):
    if align == 'c':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'r':
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == 'j':
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = spacing
    if first_indent:
        pf.first_line_indent = Pt(first_indent)
    return p


def add_text(doc, text, size=12, bold=False, cn=CN_BODY, color=None,
             align=None, first_indent=24, before=0, after=4, spacing=1.4):
    p = doc.add_paragraph()
    set_para(p, align, first_indent if align is None else 0, before, after, spacing)
    set_run(p.add_run(text), cn=cn, size=size, bold=bold, color=color)
    return p


def add_bullet(doc, text, size=11.5, level=0, cn=CN_BODY):
    p = doc.add_paragraph()
    set_para(p, None, 0, 0, 2, 1.35)
    p.paragraph_format.left_indent = Pt(24 + level * 18)
    marker = '●  ' if level == 0 else '○  '
    set_run(p.add_run(marker), cn=cn, size=size, color=BLUE)
    set_run(p.add_run(text), cn=cn, size=size)
    return p


def h1(doc, text):
    p = doc.add_paragraph()
    set_para(p, None, 0, 16, 8, 1.3)
    set_run(p.add_run(text), cn=CN_HEAD, size=15, bold=True, color=NAVY)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '8')
    bottom.set(qn('w:space'), '2')
    bottom.set(qn('w:color'), NAVY_HEX)
    pbdr.append(bottom)
    pPr.append(pbdr)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    set_para(p, None, 0, 12, 6, 1.3)
    set_run(p.add_run(text), cn=CN_HEAD, size=13, bold=True, color=BLUE)
    return p


def h3(doc, text):
    p = doc.add_paragraph()
    set_para(p, None, 0, 8, 4, 1.3)
    set_run(p.add_run(text), cn=CN_HEAD, size=12, bold=True, color=NAVY)
    return p


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)


def set_borders(table, color=BLUE_HEX, sz=4):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        borders.append(el)
    tblPr.append(borders)


def set_cell(cell, text, size=10.5, bold=False, cn=CN_BODY,
             color=None, align=None, fill=None):
    cell.text = ''
    p = cell.paragraphs[0]
    set_para(p, align, 0, 2, 2, 1.2)
    set_run(p.add_run(str(text)), cn=cn, size=size, bold=bold, color=color)
    if fill:
        shade(cell, fill)
    return cell


def add_table(doc, headers, rows, widths=None, size=10.5, head_size=10.5,
              align_cols=None, caption=None):
    """align_cols: 列索引集合 -> 居中；其余左对齐"""
    if widths:
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
    else:
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = 'Table Grid'
    for i, htxt in enumerate(headers):
        set_cell(t.rows[0].cells[i], htxt, size=head_size, bold=True,
                 cn=CN_HEAD, color=NAVY, align='c', fill=HEAD_FILL)
    for r_i, row in enumerate(rows):
        cells = t.add_row().cells
        fill = ALT_FILL if r_i % 2 == 1 else None
        for i, val in enumerate(row):
            al = 'c' if (align_cols and i in align_cols) else None
            set_cell(cells[i], val, size=size, align=al, fill=fill)
    set_borders(t)
    if widths:
        for r in t.rows:
            for i, w in enumerate(widths):
                r.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def add_caption(doc, text):
    p = doc.add_paragraph()
    set_para(p, 'c', 0, 0, 8, 1.2)
    set_run(p.add_run(text), cn=CN_BODY, size=9.5, color=GREY)
    return p


def add_cover(doc, title, subtitle, meta_rows, org_line):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    set_para(p, 'c', 0, 0, 6, 1.2)
    set_run(p.add_run(org_line), cn=CN_HEAD, size=13, color=BLUE)
    p = doc.add_paragraph()
    set_para(p, 'c', 0, 0, 18, 1.25)
    set_run(p.add_run(title), cn=CN_HEAD, size=26, bold=True, color=NAVY)
    p = doc.add_paragraph()
    set_para(p, 'c', 0, 0, 40, 1.2)
    set_run(p.add_run(subtitle), cn=CN_HEAD, size=14, color=GREY)
    t = doc.add_table(rows=len(meta_rows), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(meta_rows):
        set_cell(t.rows[i].cells[0], k, size=11, bold=True, cn=CN_HEAD,
                 color=NAVY, align='r')
        set_cell(t.rows[i].cells[1], v, size=11, align='c')
        t.rows[i].cells[0].width = Cm(4.2)
        t.rows[i].cells[1].width = Cm(7.2)
    doc.add_page_break()


def add_toc(doc, levels="1-3"):
    p = doc.add_paragraph()
    set_para(p, 'c', 0, 0, 8, 1.3)
    set_run(p.add_run("目　　录"), cn=CN_HEAD, size=16, bold=True, color=NAVY)
    p = doc.add_paragraph()
    run = p.add_run()
    set_run(run, cn=CN_BODY, size=11)
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve')
    it.text = 'TOC \\o "%s" \\h \\z \\u' % levels
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t')
    t.text = "【打开后按 Ctrl+A 全选，再按 F9 更新域，即可生成目录页码】"
    f3 = OxmlElement('w:fldChar'); f3.set(qn('w:fldCharType'), 'end')
    run._r.append(f1); run._r.append(it); run._r.append(f2)
    run._r.append(t); run._r.append(f3)
    doc.add_page_break()


def _field(paragraph, instr):
    run = paragraph.add_run()
    set_run(run, cn=CN_BODY, size=9, color=GREY)
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve')
    it.text = instr
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end')
    run._r.append(f1); run._r.append(it); run._r.append(f2)
    return run


def add_footer(doc, left_text, start_page=True):
    sec = doc.sections[0]
    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(left_text + "　　第 "), cn=CN_BODY, size=9, color=GREY)
    _field(p, 'PAGE')
    set_run(p.add_run(" 页 / 共 "), cn=CN_BODY, size=9, color=GREY)
    _field(p, 'NUMPAGES')
    set_run(p.add_run(" 页"), cn=CN_BODY, size=9, color=GREY)


def add_header(doc, text):
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(hp.add_run(text), cn=CN_BODY, size=9, color=GREY)
